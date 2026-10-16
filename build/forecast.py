import csv

import openpyxl

from layout import FORECAST_YEARS, ROOT, TARGET, YEARS

OUT = ROOT / "data" / "forecast.csv"
SUMMARY = ROOT / "data" / "forecast_summary.csv"


def find_first_year_row(worksheet):
    for row in worksheet.iter_rows(min_col=1, max_col=1):
        if row[0].value == YEARS[0]:
            return row[0].row
    raise SystemExit("Pet33-Forecast: no year rows found")


def read_forecast():
    workbook = openpyxl.load_workbook(TARGET, data_only=True)
    worksheet = workbook["Pet33-Forecast"]
    first = find_first_year_row(worksheet)
    rows = []
    for offset, year in enumerate(YEARS + FORECAST_YEARS):
        row = first + offset
        if worksheet.cell(row, 1).value != year:
            raise SystemExit(f"Pet33-Forecast row {row} is not {year}")
        rows.append(
            {
                "year": year,
                "actual": worksheet.cell(row, 2).value,
                "linear": worksheet.cell(row, 3).value,
                "cagr": worksheet.cell(row, 4).value,
                "target": worksheet.cell(row, 5).value,
                "gap": worksheet.cell(row, 6).value,
            }
        )
    summary = {}
    for row in worksheet.iter_rows(min_col=1, max_col=2):
        label, value = row[0].value, row[1].value
        if isinstance(label, str) and isinstance(value, (int, float)):
            summary[label] = value
    return rows, summary


def main():
    rows, summary = read_forecast()
    actual = [r for r in rows if r["actual"] is not None]
    if len(actual) != len(YEARS):
        raise SystemExit(f"{len(actual)} actual years, expected {len(YEARS)}")
    forecast = [r for r in rows if r["linear"] is not None]
    if len(forecast) != len(FORECAST_YEARS):
        raise SystemExit(f"{len(forecast)} forecast years, expected {len(FORECAST_YEARS)}")
    with OUT.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["year", "actual", "linear", "cagr", "target", "gap"])
        writer.writeheader()
        writer.writerows(rows)
    with SUMMARY.open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["measure", "value"])
        for label, value in summary.items():
            writer.writerow([label, value])
    print(f"wrote {OUT.relative_to(ROOT)} and {SUMMARY.relative_to(ROOT)}")
    for label in (
        "Import 2016 ($ Millions)",
        "Import 2025 ($ Millions)",
        "Change 2016 to 2025 ($ Millions)",
        "Change 2016 to 2025 (percent)",
        "Import CAGR 2016 to 2025",
        "Linear trend R squared",
        "Target import 2030 ($ Millions)",
    ):
        print(f"  {label}: {summary[label]:,.4f}")


if __name__ == "__main__":
    main()
