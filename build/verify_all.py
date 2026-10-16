import csv
import re
import subprocess
import sys
from pathlib import Path

import openpyxl

import read_workbook as rw
from layout import DIRECTIONS, FORECAST_YEARS, ROOT, TARGET, YEARS

DOCS = ROOT / "docs"
PDF = ROOT / "out" / "24503993_A3.pdf"

problems = []


def check(condition, message):
    if not condition:
        problems.append(message)
    return condition


def section(name):
    print(f"\n{name}")


def check_data():
    section("Data")
    rows = list(csv.DictReader((ROOT / "data" / "trade.csv").open()))
    check(len(rows) == 1560, f"trade.csv has {len(rows)} rows, expected 1560")
    check(
        not [r for r in rows if r["value"] in ("", None)],
        "trade.csv holds an empty value",
    )
    directions = {r["direction"] for r in rows}
    check(directions == set(DIRECTIONS), f"directions are {directions}")
    print(f"  {len(rows)} rows, both directions, no gaps")


def check_workbook():
    section("Workbook")
    result = subprocess.run(
        [sys.executable, "verify_workbook.py"],
        cwd=ROOT / "build",
        capture_output=True,
        text=True,
    )
    check(result.returncode == 0, f"verify_workbook.py failed:\n{result.stdout}\n{result.stderr}")
    print("  " + result.stdout.strip().splitlines()[-1])

    workbook = openpyxl.load_workbook(TARGET, data_only=True)
    required = [
        "README", "Keys", "Export-dataset", "Import-dataset", "Export-tidy", "Import-tidy",
        "Main-Analytics", "Main-Stats", "Balance", "Correlation", "ChangePoints",
        "Cat3-SubAnalytics", "Pet33-Focus", "Pet33-vs-Sub", "Pet33-Forecast", "AI-Comparison",
    ]
    for name in required:
        check(name in workbook.sheetnames, f"workbook is missing sheet {name}")
    charts = sum(len(getattr(workbook[n], "_charts", [])) for n in workbook.sheetnames)
    print(f"  {len(workbook.sheetnames)} sheets present")

    ai = rw.ai_comparison()
    check(all(r["ai"] is not None for r in ai), "AI-Comparison has an empty AI column")
    logged = {
        int(r["year"]): float(r["value"])
        for r in csv.DictReader((ROOT / "data" / "ai_forecast.csv").open())
    }
    for row in ai:
        check(
            logged[row["year"]] == row["ai"],
            f"AI forecast {row['year']} differs between the log and the workbook",
        )
    print(f"  AI comparison matches the logged run for all {len(ai)} years")


def figure_titles(html):
    return re.findall(r'"title":\s*\{(?:[^{}]|\{[^{}]*\})*?"text":\s*"([^"]+)"', html)


def check_pages():
    section("Web pages")
    for name in ("index.html", "dashboard.html", "storyboard.html"):
        path = DOCS / name
        check(path.exists(), f"{name} is missing")
        if not path.exists():
            continue
        html = path.read_text()
        loaders = [
            tag
            for tag in re.findall(r"<(?:script|link|img|iframe|source)\b[^>]*>", html, re.I)
            if re.search(r"\b(?:src|href)\s*=", tag)
        ]
        check(not loaders, f"{name} loads an external file: {loaders[:2]}")
        lowered = html.lower()
        for direction in DIRECTIONS:
            check(
                direction.lower() in lowered,
                f"{name} does not mention {direction}",
            )
        print(f"  {name}: self contained, both directions, {len(html) / 1024 / 1024:.1f} MB")

    dashboard = (DOCS / "dashboard.html").read_text()
    comparisons = len(re.findall(r"Comparison \d\.", dashboard))
    check(comparisons >= 5, f"dashboard has {comparisons} comparisons, expected 5")
    insights = len(re.findall(r'<span class="yr">', dashboard))
    check(insights == 10, f"dashboard has {insights} insight lines, expected 10")
    tables = len(re.findall(r'<table class="data"', dashboard))
    check(tables >= 3, f"dashboard has {tables} tables, expected at least 3")
    print(f"  dashboard: {comparisons} comparisons, {insights} year insights, {tables} tables")

    storyboard = (DOCS / "storyboard.html").read_text()
    panels = len(re.findall(r'<section class="panel">', storyboard))
    check(panels >= 12, f"storyboard has {panels} panels, expected at least 12")
    print(f"  storyboard: {panels} panels")

    for name in ("dashboard.html", "storyboard.html"):
        html = (DOCS / name).read_text()
        titles = [t for t in figure_titles(html) if len(t) > 2]
        check(len(titles) >= 8, f"{name} has only {len(titles)} chart or axis titles")


def check_colours():
    section("Colour")
    import theme

    for name in ("dashboard.html", "storyboard.html", "index.html"):
        html = (DOCS / name).read_text()
        check(theme.EXPORT.lower() in html.lower(), f"{name} does not use the export colour")
        check(theme.IMPORT.lower() in html.lower(), f"{name} does not use the import colour")
    print(f"  export {theme.EXPORT} and import {theme.IMPORT} used on every page")


def pdf_text():
    result = subprocess.run(
        ["pdftotext", str(PDF), "-"], capture_output=True, text=True
    )
    return result.stdout


def check_report():
    section("Report")
    check(PDF.exists(), "the PDF is missing")
    if not PDF.exists():
        return
    info = subprocess.run(["pdfinfo", str(PDF)], capture_output=True, text=True).stdout
    pages = int(re.search(r"Pages:\s+(\d+)", info).group(1))
    check(pages <= 20, f"the report is {pages} pages, over the 20 page limit")
    check("A4" in info, "the report is not A4")
    text = pdf_text()
    figures = sorted(set(int(n) for n in re.findall(r"Figure (\d+)\.", text)))
    check(figures == list(range(1, len(figures) + 1)), f"figure numbering is {figures}")
    check(len(figures) >= 12, f"the report has {len(figures)} figures, expected at least 12")
    print(f"  {pages} pages, A4, {len(figures)} figures numbered 1 to {len(figures)}")

    for heading in (
        "Executive summary",
        "Data collection and preparation",
        "Dataset characteristics",
        "All ten main categories",
        "change point in each year",
        "Main category against subcategory",
        "The dashboard",
        "The storyboard",
        "Trends and outliers",
        "the target, and the forecast",
        "Comparison against the Claude Sonnet 5 forecast",
        "Visualisation techniques",
        "The dashboard against the storyboard",
        "Recommendations",
        "Conclusion",
        "The Claude Sonnet 5 prompt and the reply",
    ):
        check(heading in text, f"the report is missing the section {heading!r}")
    print("  all 16 required sections present")

    summary = rw.forecast_summary()
    focus = rw.focus()
    ai = rw.ai_comparison()
    expected = {
        "2016 petroleum import": f"{focus[0]['import']:,.0f}",
        "2025 petroleum import": f"{focus[-1]['import']:,.0f}",
        "2025 petroleum export": f"{focus[-1]['export']:,.0f}",
        "2016 petroleum deficit": f"{abs(focus[0]['balance']):,.0f}",
        "2025 petroleum deficit": f"{abs(focus[-1]['balance']):,.0f}",
        "ten year change": f"{summary['Change 2016 to 2025 ($ Millions)']:,.0f}",
        "2030 target": f"{summary['Target import 2030 ($ Millions)']:,.0f}",
        "AI 2030 forecast": f"{ai[-1]['ai']:,.0f}",
    }
    flat = text.replace("\n", " ")
    for label, value in expected.items():
        check(value in flat, f"the report does not quote the {label} of {value}")
    print(f"  all {len(expected)} headline numbers match the workbook")

    check("REPLACE-ME" not in text, "the report still holds the placeholder Pages URL")


def main():
    check_data()
    check_workbook()
    check_pages()
    check_colours()
    check_report()
    print()
    if problems:
        print(f"FAILED with {len(problems)} problems")
        for message in problems:
            print("  -", message)
        raise SystemExit(1)
    print("All checks passed.")


if __name__ == "__main__":
    main()
