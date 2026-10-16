import openpyxl

from layout import DIRECTIONS, FORECAST_YEARS, TARGET, TOTAL_CODE, YEARS

FOCUS_MAIN = "3"
FOCUS_SUB = "33"
FOCUS_SUBS = ["32", "33", "34"]

_cache = {}


def workbook():
    if "wb" not in _cache:
        _cache["wb"] = openpyxl.load_workbook(TARGET, data_only=True)
    return _cache["wb"]


def find_title(worksheet, title):
    for row in worksheet.iter_rows(min_col=1, max_col=1):
        if row[0].value == title:
            return row[0].row
    raise SystemExit(f"{worksheet.title}: no block titled {title!r}")


def block(sheet_name, title, width):
    worksheet = workbook()[sheet_name]
    header = find_title(worksheet, title)
    first = header + 2
    names = {}
    values = {}
    row = first
    while worksheet.cell(row, 1).value not in (None, ""):
        code = str(worksheet.cell(row, 1).value)
        names[code] = worksheet.cell(row, 2).value
        values[code] = [worksheet.cell(row, 3 + i).value for i in range(width)]
        row += 1
    return names, values


def main_values(direction):
    return block("Main-Analytics", f"{direction}: value ($ Millions)", len(YEARS))


def main_growth(direction):
    return block("Main-Analytics", f"{direction}: year on year growth", len(YEARS) - 1)


def main_index(direction):
    return block("Main-Analytics", f"{direction}: index (2016 = 100)", len(YEARS))


def main_share(direction):
    return block(
        "Main-Analytics", f"{direction}: share of {direction.lower()} total", len(YEARS)
    )


def main_summary(direction):
    return block("Main-Analytics", f"{direction}: decade summary", 3)


def main_stats(direction):
    return block(
        "Main-Stats", f"{direction}: descriptive and trend statistics ($ Millions)", 10
    )


def main_zscore(direction):
    return block(
        "Main-Stats", f"{direction}: outlier z-score of year on year growth", len(YEARS) - 1
    )


def balance():
    return block("Balance", "Trade balance, export minus import ($ Millions)", len(YEARS))


def cover():
    return block("Balance", "Cover ratio, export divided by import", len(YEARS))


def correlation(direction):
    worksheet = workbook()["Correlation"]
    header = find_title(worksheet, f"{direction}: growth correlation matrix")
    first = header + 2
    codes = [str(c) for c in range(10)]
    matrix = {}
    for i, code in enumerate(codes):
        matrix[code] = {
            other: worksheet.cell(first + i, 3 + j).value
            for j, other in enumerate(codes)
        }
    return codes, matrix


def change_points():
    worksheet = workbook()["ChangePoints"]
    header = find_title(worksheet, "Summary of change points")
    first = header + 2
    rows = []
    for offset, year in enumerate(YEARS[1:]):
        row = first + offset
        rows.append(
            {
                "year": year,
                "export_rise_name": worksheet.cell(row, 2).value,
                "export_rise": worksheet.cell(row, 3).value,
                "export_fall_name": worksheet.cell(row, 4).value,
                "export_fall": worksheet.cell(row, 5).value,
                "import_rise_name": worksheet.cell(row, 6).value,
                "import_rise": worksheet.cell(row, 7).value,
                "import_fall_name": worksheet.cell(row, 8).value,
                "import_fall": worksheet.cell(row, 9).value,
                "cause": worksheet.cell(row, 10).value,
            }
        )
    return rows


def cat3_values(direction):
    return block("Cat3-SubAnalytics", f"{direction}: value ($ Millions)", len(YEARS))


def cat3_share(direction):
    return block("Cat3-SubAnalytics", f"{direction}: share of main category 3", len(YEARS))


def cat3_growth(direction):
    return block("Cat3-SubAnalytics", f"{direction}: year on year growth", len(YEARS) - 1)


def cat3_index(direction):
    return block("Cat3-SubAnalytics", f"{direction}: index (2016 = 100)", len(YEARS))


def cat3_correl(direction):
    return block(
        "Cat3-SubAnalytics",
        f"{direction}: growth correlation against main category 3",
        1,
    )


def focus():
    worksheet = workbook()["Pet33-Focus"]
    first = None
    for row in worksheet.iter_rows(min_col=1, max_col=1):
        if row[0].value == YEARS[0]:
            first = row[0].row
            break
    rows = []
    for offset, year in enumerate(YEARS):
        row = first + offset
        rows.append(
            {
                "year": year,
                "export": worksheet.cell(row, 2).value,
                "import": worksheet.cell(row, 3).value,
                "balance": worksheet.cell(row, 4).value,
                "cover": worksheet.cell(row, 5).value,
                "export_growth": worksheet.cell(row, 6).value,
                "import_growth": worksheet.cell(row, 7).value,
                "export_ma": worksheet.cell(row, 8).value,
                "import_ma": worksheet.cell(row, 9).value,
                "export_z": worksheet.cell(row, 10).value,
                "import_z": worksheet.cell(row, 11).value,
                "insight": worksheet.cell(row, 12).value,
            }
        )
    return rows


def vs_index(direction):
    return block("Pet33-vs-Sub", f"{direction}: index (2016 = 100)", len(YEARS))


def vs_growth(direction):
    return block("Pet33-vs-Sub", f"{direction}: year on year growth", len(YEARS) - 1)


def vs_pairs():
    worksheet = workbook()["Pet33-vs-Sub"]
    header = find_title(worksheet, "Pairwise growth correlation")
    first = header + 2
    pairs = {}
    row = first
    while worksheet.cell(row, 1).value not in (None, ""):
        label = worksheet.cell(row, 1).value
        pairs[label] = {
            "Export": worksheet.cell(row, 2).value,
            "Import": worksheet.cell(row, 3).value,
        }
        row += 1
    return pairs, worksheet.cell(first, 4).value


def forecast():
    worksheet = workbook()["Pet33-Forecast"]
    first = None
    for row in worksheet.iter_rows(min_col=1, max_col=1):
        if row[0].value == YEARS[0]:
            first = row[0].row
            break
    rows = []
    for offset, year in enumerate(YEARS + FORECAST_YEARS):
        row = first + offset
        rows.append(
            {
                "year": year,
                "actual": worksheet.cell(row, 2).value,
                "linear": worksheet.cell(row, 3).value,
                "cagr": worksheet.cell(row, 4).value,
                "target": worksheet.cell(row, 5).value,
                "gap": worksheet.cell(row, 6).value,
            }
        )
    return rows


def forecast_summary():
    worksheet = workbook()["Pet33-Forecast"]
    out = {}
    for row in worksheet.iter_rows(min_col=1, max_col=2):
        label, value = row[0].value, row[1].value
        if isinstance(label, str) and isinstance(value, (int, float)):
            out[label] = value
    return out


def ai_comparison():
    worksheet = workbook()["AI-Comparison"]
    first = None
    for row in worksheet.iter_rows(min_col=1, max_col=1):
        if row[0].value == FORECAST_YEARS[0]:
            first = row[0].row
            break
    rows = []
    for offset, year in enumerate(FORECAST_YEARS):
        row = first + offset
        rows.append(
            {
                "year": year,
                "linear": worksheet.cell(row, 2).value,
                "cagr": worksheet.cell(row, 3).value,
                "ai": worksheet.cell(row, 4).value,
                "difference": worksheet.cell(row, 5).value,
                "percent": worksheet.cell(row, 6).value,
            }
        )
    return rows


def totals(direction):
    worksheet = workbook()[f"{direction}-dataset"]
    header = [c.value for c in worksheet[1]]
    column = None
    for index, value in enumerate(header):
        if isinstance(value, str) and value.strip().startswith(TOTAL_CODE):
            column = index + 1
            break
    return {
        YEARS[i]: worksheet.cell(4 + i, column).value for i in range(len(YEARS))
    }


def backtest():
    worksheet = workbook()["Pet33-Backtest"]
    first = None
    for row in worksheet.iter_rows(min_col=1, max_col=1):
        if row[0].value == YEARS[0]:
            first = row[0].row
            break
    rows = []
    for offset, year in enumerate(YEARS):
        row = first + offset
        rows.append(
            {
                "year": year,
                "actual": worksheet.cell(row, 2).value,
                "linear": worksheet.cell(row, 3).value,
                "cagr": worksheet.cell(row, 4).value,
                "recent": worksheet.cell(row, 5).value,
                "linear_error": worksheet.cell(row, 6).value,
                "cagr_error": worksheet.cell(row, 7).value,
                "recent_error": worksheet.cell(row, 8).value,
            }
        )
    mape_row = first + len(YEARS) + 1
    mape = {
        "linear": worksheet.cell(mape_row, 6).value,
        "cagr": worksheet.cell(mape_row, 7).value,
        "recent": worksheet.cell(mape_row, 8).value,
    }
    return rows, mape


def sheet_count():
    return len(workbook().sheetnames)
