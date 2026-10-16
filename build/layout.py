from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "2026-Spring-Ass3-V20.xlsx"
TARGET = ROOT / "out" / "24503993_A3.xlsx"

YEARS = list(range(2016, 2026))
FORECAST_YEARS = list(range(2026, 2031))
DIRECTIONS = ["Export", "Import"]
TOTAL_CODE = "-1"
FOCUS_MAIN = "3"
FOCUS_SUBS = ["32", "33", "34"]
FOCUS_SUB = "33"

HEADER_ROW = 1
UNIT_ROW = 2
SERIES_ROW = 3
FIRST_YEAR_ROW = 4


def clean(raw):
    text = str(raw).strip()
    if text.endswith(";"):
        text = text[:-1]
    return text.strip()


def split_header(raw):
    code, _, name = clean(raw).partition(" ")
    return code.strip(), name.strip()


def level_of(code):
    if code == TOTAL_CODE:
        return "total"
    return "main" if len(code) == 1 else "sub"


def read_source():
    workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    layout = {}
    for direction in DIRECTIONS:
        worksheet = workbook[f"{direction}-dataset"]
        rows = list(worksheet.iter_rows(values_only=True))
        columns = {}
        names = {}
        order = []
        for index in range(1, len(rows[0])):
            code, name = split_header(rows[0][index])
            columns[code] = get_column_letter(index + 1)
            names[code] = name
            order.append(code)
        values = {}
        for row in rows[FIRST_YEAR_ROW - 1 :]:
            year = int(str(row[0]).lstrip("Y"))
            for index in range(1, len(rows[0])):
                code, _ = split_header(rows[0][index])
                values.setdefault(code, {})[year] = float(row[index])
        layout[direction] = {
            "rows": rows,
            "columns": columns,
            "names": names,
            "order": order,
            "values": values,
            "mains": [c for c in order if level_of(c) == "main"],
            "subs": [c for c in order if level_of(c) == "sub"],
        }
    return layout


def source_ref(direction, code, year, layout):
    column = layout[direction]["columns"][code]
    row = FIRST_YEAR_ROW + (year - YEARS[0])
    return f"'{direction}-dataset'!{column}{row}"
