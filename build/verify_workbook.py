import csv
import math
from collections import defaultdict

import openpyxl

from layout import DIRECTIONS, FORECAST_YEARS, ROOT, TARGET, TOTAL_CODE, YEARS

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!", "Err:")

FOCUS_MAIN = "3"
FOCUS_SUB = "33"
FOCUS_SUBS = ["32", "33", "34"]

failures = []


def fail(message):
    failures.append(message)


def load_values():
    rows = list(csv.DictReader((ROOT / "data" / "trade.csv").open()))
    table = defaultdict(dict)
    for row in rows:
        table[(row["direction"], row["code"])][int(row["year"])] = float(row["value"])
    return table


def close(left, right, tolerance=1e-6):
    if left is None or right is None:
        return False
    if isinstance(left, str) or isinstance(right, str):
        return False
    return abs(left - right) <= tolerance * max(1.0, abs(right))


def scan_for_errors(workbook):
    for name in workbook.sheetnames:
        worksheet = workbook[name]
        for row in worksheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and any(value.startswith(e) or value == e for e in ERRORS):
                    fail(f"{name}!{cell.coordinate} holds the error {value}")


def find_block(worksheet, title):
    for row in worksheet.iter_rows(min_col=1, max_col=1):
        if row[0].value == title:
            return row[0].row
    fail(f"{worksheet.title}: block titled {title!r} not found")
    return None


def read_block(worksheet, title, width):
    header = find_block(worksheet, title)
    if header is None:
        return {}
    first = header + 2
    out = {}
    row = first
    while worksheet.cell(row, 1).value not in (None, ""):
        code = str(worksheet.cell(row, 1).value)
        out[code] = [worksheet.cell(row, 3 + i).value for i in range(width)]
        row += 1
    return out


def check_main_analytics(workbook, table):
    worksheet = workbook["Main-Analytics"]
    for direction in DIRECTIONS:
        values = read_block(worksheet, f"{direction}: value ($ Millions)", len(YEARS))
        growth = read_block(worksheet, f"{direction}: year on year growth", len(YEARS) - 1)
        index = read_block(worksheet, f"{direction}: index (2016 = 100)", len(YEARS))
        share = read_block(worksheet, f"{direction}: share of {direction.lower()} total", len(YEARS))
        if len(values) != 10:
            fail(f"Main-Analytics {direction}: {len(values)} value rows, expected 10")
        for code, row in values.items():
            for offset, year in enumerate(YEARS):
                if not close(row[offset], table[(direction, code)][year]):
                    fail(f"Main-Analytics {direction} value {code} {year}: {row[offset]}")
        for code, row in growth.items():
            for offset, year in enumerate(YEARS[1:]):
                expected = table[(direction, code)][year] / table[(direction, code)][year - 1] - 1
                if not close(row[offset], expected):
                    fail(f"Main-Analytics {direction} growth {code} {year}: {row[offset]} vs {expected}")
        for code, row in index.items():
            for offset, year in enumerate(YEARS):
                expected = table[(direction, code)][year] / table[(direction, code)][YEARS[0]] * 100
                if not close(row[offset], expected):
                    fail(f"Main-Analytics {direction} index {code} {year}: {row[offset]} vs {expected}")
        for code, row in share.items():
            for offset, year in enumerate(YEARS):
                expected = table[(direction, code)][year] / table[(direction, TOTAL_CODE)][year]
                if not close(row[offset], expected):
                    fail(f"Main-Analytics {direction} share {code} {year}: {row[offset]} vs {expected}")
        summary = read_block(worksheet, f"{direction}: decade summary", 3)
        for code, row in summary.items():
            series = table[(direction, code)]
            span = YEARS[-1] - YEARS[0]
            cagr = (series[YEARS[-1]] / series[YEARS[0]]) ** (1 / span) - 1
            contribution = series[YEARS[-1]] - series[YEARS[0]]
            totals = table[(direction, TOTAL_CODE)]
            growth_share = contribution / (totals[YEARS[-1]] - totals[YEARS[0]])
            for got, expected, label in zip(row, [cagr, contribution, growth_share], ["cagr", "contribution", "growth share"]):
                if not close(got, expected):
                    fail(f"Main-Analytics {direction} {label} {code}: {got} vs {expected}")


def stdevp(values):
    mean = sum(values) / len(values)
    return math.sqrt(sum((v - mean) ** 2 for v in values) / len(values))


def linear_fit(xs, ys):
    n = len(xs)
    mean_x = sum(xs) / n
    mean_y = sum(ys) / n
    sxx = sum((x - mean_x) ** 2 for x in xs)
    sxy = sum((x - mean_x) * (y - mean_y) for x, y in zip(xs, ys))
    slope = sxy / sxx
    intercept = mean_y - slope * mean_x
    syy = sum((y - mean_y) ** 2 for y in ys)
    rsq = (sxy ** 2) / (sxx * syy)
    return slope, intercept, rsq


def pearson(xs, ys):
    n = len(xs)
    mean_x = sum(xs) / n
    mean_y = sum(ys) / n
    sxy = sum((x - mean_x) * (y - mean_y) for x, y in zip(xs, ys))
    sxx = sum((x - mean_x) ** 2 for x in xs)
    syy = sum((y - mean_y) ** 2 for y in ys)
    return sxy / math.sqrt(sxx * syy)


def check_main_stats(workbook, table):
    worksheet = workbook["Main-Stats"]
    for direction in DIRECTIONS:
        stats = read_block(
            worksheet, f"{direction}: descriptive and trend statistics ($ Millions)", 10
        )
        if len(stats) != 10:
            fail(f"Main-Stats {direction}: {len(stats)} rows, expected 10")
        for code, row in stats.items():
            series = [table[(direction, code)][y] for y in YEARS]
            ordered = sorted(series)
            median = (ordered[4] + ordered[5]) / 2
            slope, intercept, rsq = linear_fit(YEARS, series)
            expected = [
                sum(series) / len(series),
                median,
                stdevp(series),
                stdevp(series) / (sum(series) / len(series)),
                min(series),
                max(series),
                max(series) - min(series),
                slope,
                intercept,
                rsq,
            ]
            labels = ["mean", "median", "sd", "cv", "min", "max", "range", "slope", "intercept", "rsq"]
            for got, want, label in zip(row, expected, labels):
                if not close(got, want):
                    fail(f"Main-Stats {direction} {label} {code}: {got} vs {want}")
        zscores = read_block(
            worksheet, f"{direction}: outlier z-score of year on year growth", len(YEARS) - 1
        )
        for code, row in zscores.items():
            series = table[(direction, code)]
            growth = [series[y] / series[y - 1] - 1 for y in YEARS[1:]]
            mean = sum(growth) / len(growth)
            spread = stdevp(growth)
            for offset, value in enumerate(growth):
                expected = (value - mean) / spread
                if not close(row[offset], expected):
                    fail(f"Main-Stats {direction} z {code} {YEARS[1:][offset]}: {row[offset]} vs {expected}")


def check_balance(workbook, table):
    worksheet = workbook["Balance"]
    balance = read_block(worksheet, "Trade balance, export minus import ($ Millions)", len(YEARS))
    cover = read_block(worksheet, "Cover ratio, export divided by import", len(YEARS))
    if len(balance) != 11:
        fail(f"Balance: {len(balance)} rows, expected 11")
    for code, row in balance.items():
        for offset, year in enumerate(YEARS):
            expected = table[("Export", code)][year] - table[("Import", code)][year]
            if not close(row[offset], expected):
                fail(f"Balance {code} {year}: {row[offset]} vs {expected}")
    for code, row in cover.items():
        for offset, year in enumerate(YEARS):
            expected = table[("Export", code)][year] / table[("Import", code)][year]
            if not close(row[offset], expected):
                fail(f"Cover {code} {year}: {row[offset]} vs {expected}")


def check_correlation(workbook, table):
    worksheet = workbook["Correlation"]
    mains = [str(c) for c in range(10)]
    for direction in DIRECTIONS:
        header = find_block(worksheet, f"{direction}: growth correlation matrix")
        if header is None:
            continue
        first = header + 2
        for i, code in enumerate(mains):
            for j, other in enumerate(mains):
                got = worksheet.cell(first + i, 3 + j).value
                left = table[(direction, code)]
                right = table[(direction, other)]
                gl = [left[y] / left[y - 1] - 1 for y in YEARS[1:]]
                gr = [right[y] / right[y - 1] - 1 for y in YEARS[1:]]
                if not close(got, pearson(gl, gr)):
                    fail(f"Correlation {direction} {code}x{other}: {got} vs {pearson(gl, gr)}")


def check_change_points(workbook, table):
    worksheet = workbook["ChangePoints"]
    header = find_block(worksheet, "Summary of change points")
    if header is None:
        return
    first = header + 2
    count = 0
    for offset, year in enumerate(YEARS[1:]):
        row = first + offset
        if worksheet.cell(row, 1).value != year:
            fail(f"ChangePoints row {row}: year {worksheet.cell(row, 1).value} not {year}")
            continue
        count += 1
        column = 2
        for direction in DIRECTIONS:
            changes = {
                code: table[(direction, code)][year] - table[(direction, code)][year - 1]
                for code in [str(c) for c in range(10)]
            }
            top = max(changes.values())
            bottom = min(changes.values())
            if not close(worksheet.cell(row, column + 1).value, top):
                fail(f"ChangePoints {year} {direction} rise: {worksheet.cell(row, column + 1).value} vs {top}")
            if not close(worksheet.cell(row, column + 3).value, bottom):
                fail(f"ChangePoints {year} {direction} fall: {worksheet.cell(row, column + 3).value} vs {bottom}")
            column += 4
        if not worksheet.cell(row, 10).value:
            fail(f"ChangePoints {year}: no cause text")
    if count != 9:
        fail(f"ChangePoints: {count} year rows, expected 9")


def check_focus(workbook, table):
    worksheet = workbook["Pet33-Focus"]
    first = None
    for row in worksheet.iter_rows(min_col=1, max_col=1):
        if row[0].value == YEARS[0]:
            first = row[0].row
            break
    if first is None:
        fail("Pet33-Focus: no year rows")
        return
    insights = 0
    for offset, year in enumerate(YEARS):
        row = first + offset
        export = table[("Export", FOCUS_SUB)][year]
        imports = table[("Import", FOCUS_SUB)][year]
        for column, expected, label in (
            (2, export, "export"),
            (3, imports, "import"),
            (4, export - imports, "balance"),
            (5, export / imports, "cover"),
        ):
            if not close(worksheet.cell(row, column).value, expected):
                fail(f"Pet33-Focus {year} {label}: {worksheet.cell(row, column).value} vs {expected}")
        if worksheet.cell(row, 12).value:
            insights += 1
    if insights != 10:
        fail(f"Pet33-Focus: {insights} insight lines, expected 10")


def check_cat3(workbook, table):
    worksheet = workbook["Cat3-SubAnalytics"]
    codes = [FOCUS_MAIN] + FOCUS_SUBS
    for direction in DIRECTIONS:
        values = read_block(worksheet, f"{direction}: value ($ Millions)", len(YEARS))
        share = read_block(worksheet, f"{direction}: share of main category 3", len(YEARS))
        growth = read_block(worksheet, f"{direction}: year on year growth", len(YEARS) - 1)
        index = read_block(worksheet, f"{direction}: index (2016 = 100)", len(YEARS))
        if sorted(values) != sorted(codes):
            fail(f"Cat3 {direction}: codes {sorted(values)} not {sorted(codes)}")
        for code in codes:
            series = table[(direction, code)]
            main = table[(direction, FOCUS_MAIN)]
            for offset, year in enumerate(YEARS):
                if not close(values[code][offset], series[year]):
                    fail(f"Cat3 {direction} value {code} {year}")
                if not close(share[code][offset], series[year] / main[year]):
                    fail(f"Cat3 {direction} share {code} {year}")
                if not close(index[code][offset], series[year] / series[YEARS[0]] * 100):
                    fail(f"Cat3 {direction} index {code} {year}")
            for offset, year in enumerate(YEARS[1:]):
                if not close(growth[code][offset], series[year] / series[year - 1] - 1):
                    fail(f"Cat3 {direction} growth {code} {year}")
        correl = read_block(
            worksheet, f"{direction}: growth correlation against main category 3", 1
        )
        main_growth = [
            table[(direction, FOCUS_MAIN)][y] / table[(direction, FOCUS_MAIN)][y - 1] - 1
            for y in YEARS[1:]
        ]
        for code in codes:
            series = table[(direction, code)]
            own = [series[y] / series[y - 1] - 1 for y in YEARS[1:]]
            if not close(correl[code][0], pearson(own, main_growth)):
                fail(f"Cat3 {direction} correlation {code}: {correl[code][0]}")


def check_vs_sub(workbook, table):
    worksheet = workbook["Pet33-vs-Sub"]
    for direction in DIRECTIONS:
        index = read_block(worksheet, f"{direction}: index (2016 = 100)", len(YEARS))
        growth = read_block(worksheet, f"{direction}: year on year growth", len(YEARS) - 1)
        if sorted(index) != sorted(FOCUS_SUBS):
            fail(f"Pet33-vs-Sub {direction}: codes {sorted(index)}")
        for code in FOCUS_SUBS:
            series = table[(direction, code)]
            for offset, year in enumerate(YEARS):
                if not close(index[code][offset], series[year] / series[YEARS[0]] * 100):
                    fail(f"Pet33-vs-Sub {direction} index {code} {year}")
            for offset, year in enumerate(YEARS[1:]):
                if not close(growth[code][offset], series[year] / series[year - 1] - 1):
                    fail(f"Pet33-vs-Sub {direction} growth {code} {year}")
    header = find_block(worksheet, "Pairwise growth correlation")
    if header is None:
        return
    first = header + 2

    def growth_of(direction, code):
        series = table[(direction, code)]
        return [series[y] / series[y - 1] - 1 for y in YEARS[1:]]

    pairs = [(FOCUS_SUB, "32"), (FOCUS_SUB, "34"), ("32", "34")]
    for offset, (left, right) in enumerate(pairs):
        row = first + offset
        if worksheet.cell(row, 1).value != f"{left} against {right}":
            fail(f"Pet33-vs-Sub pair row {row}: {worksheet.cell(row, 1).value}")
            continue
        for column, direction in enumerate(DIRECTIONS, start=2):
            expected = pearson(growth_of(direction, left), growth_of(direction, right))
            if not close(worksheet.cell(row, column).value, expected):
                fail(f"Pet33-vs-Sub {direction} {left}x{right}: {worksheet.cell(row, column).value} vs {expected}")
    expected = pearson(growth_of("Export", FOCUS_SUB), growth_of("Import", FOCUS_SUB))
    got = worksheet.cell(first, 4).value
    if not close(got, expected):
        fail(f"Pet33-vs-Sub import against export: {got} vs {expected}")


def check_forecast(workbook, table):
    worksheet = workbook["Pet33-Forecast"]
    first = None
    for row in worksheet.iter_rows(min_col=1, max_col=1):
        if row[0].value == YEARS[0]:
            first = row[0].row
            break
    if first is None:
        fail("Pet33-Forecast: no year rows")
        return
    series = [table[("Import", FOCUS_SUB)][y] for y in YEARS]
    slope, intercept, _ = linear_fit(YEARS, series)
    span = YEARS[-1] - YEARS[0]
    rate = (series[-1] / series[0]) ** (1 / span)
    forecast_first = first + len(YEARS)
    for offset, year in enumerate(FORECAST_YEARS):
        row = forecast_first + offset
        if worksheet.cell(row, 1).value != year:
            fail(f"Pet33-Forecast row {row}: year {worksheet.cell(row, 1).value} not {year}")
            continue
        linear = slope * year + intercept
        cagr = series[-1] * rate ** (year - YEARS[-1])
        target = series[-1] * (1 - 0.03) ** (year - YEARS[-1])
        for column, expected, label in (
            (3, linear, "linear"),
            (4, cagr, "cagr"),
            (5, target, "target"),
            (6, linear - target, "gap"),
        ):
            if not close(worksheet.cell(row, column).value, expected, 1e-5):
                fail(f"Pet33-Forecast {year} {label}: {worksheet.cell(row, column).value} vs {expected}")


def check_backtest(workbook, table):
    worksheet = workbook["Pet33-Backtest"]
    first = None
    for row in worksheet.iter_rows(min_col=1, max_col=1):
        if row[0].value == YEARS[0]:
            first = row[0].row
            break
    if first is None:
        fail("Pet33-Backtest: no year rows")
        return
    series = table[("Import", FOCUS_SUB)]
    cut = 2021
    train_years = [y for y in YEARS if y <= cut]
    train = [series[y] for y in train_years]
    slope, intercept, _ = linear_fit(train_years, train)
    rate = (train[-1] / train[0]) ** (1 / (len(train_years) - 1))
    step = (train[-1] - train[-4]) / 3
    errors = {"linear": [], "cagr": [], "recent": []}
    for offset, year in enumerate(YEARS):
        if year <= cut:
            continue
        row = first + offset
        ahead = year - cut
        expected = {
            3: slope * year + intercept,
            4: train[-1] * rate ** ahead,
            5: train[-1] + step * ahead,
        }
        for column, want in expected.items():
            got = worksheet.cell(row, column).value
            if not close(got, want, 1e-5):
                fail(f"Pet33-Backtest {year} column {column}: {got} vs {want}")
        for column, key in ((6, "linear"), (7, "cagr"), (8, "recent")):
            want = abs(expected[column - 3] - series[year]) / series[year]
            got = worksheet.cell(row, column).value
            if not close(got, want, 1e-5):
                fail(f"Pet33-Backtest {year} error column {column}: {got} vs {want}")
            errors[key].append(want)
    mape_row = first + len(YEARS) + 1
    for column, key in ((6, "linear"), (7, "cagr"), (8, "recent")):
        want = sum(errors[key]) / len(errors[key])
        got = worksheet.cell(mape_row, column).value
        if not close(got, want, 1e-5):
            fail(f"Pet33-Backtest mean error {key}: {got} vs {want}")


def check_charts(workbook):
    total = 0
    for name in workbook.sheetnames:
        worksheet = workbook[name]
        for chart in getattr(worksheet, "_charts", []):
            total += 1
            label = f"{name} chart {total}"
            if not chart_text(chart.title):
                fail(f"{label}: no title")
            if not chart_text(chart.x_axis.title):
                fail(f"{label}: no x axis title")
            if not chart_text(chart.y_axis.title):
                fail(f"{label}: no y axis title")
    if total < 9:
        fail(f"only {total} charts, expected at least 9")
    return total


def chart_text(title):
    if title is None:
        return ""
    rich = getattr(title, "tx", None)
    if rich is None:
        return str(title)
    body = getattr(rich, "rich", None)
    if body is None:
        return ""
    out = []
    for paragraph in body.p:
        for run in paragraph.r or []:
            out.append(run.t or "")
    return "".join(out).strip()


def check_both_directions(workbook):
    required = [
        "Main-Analytics",
        "Main-Stats",
        "Correlation",
        "ChangePoints",
        "Cat3-SubAnalytics",
        "Pet33-vs-Sub",
    ]
    for name in required:
        text = []
        for row in workbook[name].iter_rows(min_col=1, max_col=2):
            for cell in row:
                if isinstance(cell.value, str):
                    text.append(cell.value)
        joined = " ".join(text)
        for direction in DIRECTIONS:
            if direction not in joined:
                fail(f"{name}: no {direction} block")


def main():
    table = load_values()
    workbook = openpyxl.load_workbook(TARGET, data_only=True)
    scan_for_errors(workbook)
    check_main_analytics(workbook, table)
    check_main_stats(workbook, table)
    check_balance(workbook, table)
    check_correlation(workbook, table)
    check_change_points(workbook, table)
    check_focus(workbook, table)
    check_cat3(workbook, table)
    check_vs_sub(workbook, table)
    check_forecast(workbook, table)
    check_backtest(workbook, table)
    check_both_directions(workbook)
    charts = openpyxl.load_workbook(TARGET)
    count = check_charts(charts)
    if failures:
        print(f"{len(failures)} problems found")
        for message in failures[:40]:
            print("  -", message)
        raise SystemExit(1)
    print(f"workbook verified: {len(workbook.sheetnames)} sheets, {count} charts, no problems")


if __name__ == "__main__":
    main()
